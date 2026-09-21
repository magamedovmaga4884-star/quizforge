import os
import logging
from pathlib import Path
from typing import Tuple

logger = logging.getLogger(__name__)

class DocumentParser:
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt"}

    @classmethod
    def is_supported(cls, filename: str) -> bool:
        ext = Path(filename).suffix.lower()
        return ext in cls.SUPPORTED_EXTENSIONS

    @classmethod
    def extract_text(cls, file_path: str, filename: str) -> Tuple[bool, str]:
        """
        Extract text from file. Returns (success: bool, text_or_error: str).
        """
        if not os.path.exists(file_path):
            return False, "Файл не найден на сервере."

        ext = Path(filename).suffix.lower()

        try:
            if ext == ".txt":
                return cls._extract_txt(file_path)
            elif ext == ".pdf":
                return cls._extract_pdf(file_path)
            elif ext == ".docx":
                return cls._extract_docx(file_path)
            elif ext == ".pptx":
                return cls._extract_pptx(file_path)
            else:
                return False, f"Этот формат ({ext}) пока не поддерживается."
        except Exception as e:
            logger.error(f"Error parsing document {filename}: {e}", exc_info=True)
            return False, f"Ошибка при обработке файла {filename}: {str(e)}"

    @classmethod
    def _extract_txt(cls, file_path: str) -> Tuple[bool, str]:
        encodings = ["utf-8", "cp1251", "latin-1"]
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    text = f.read().strip()
                if text:
                    return True, text
            except (UnicodeDecodeError, Exception):
                continue
        return False, "Не удалось прочитать текстовый файл в поддерживаемой кодировке."

    @classmethod
    def _extract_pdf(cls, file_path: str) -> Tuple[bool, str]:
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            pages_text = []
            for idx, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    pages_text.append(page_text.strip())
            
            full_text = "\n\n".join(pages_text).strip()
            if not full_text:
                return False, "PDF файл не содержит распознаваемого текста или является сканированным изображением."
            return True, full_text
        except Exception as e:
            logger.warning(f"pypdf error: {e}")
            return False, f"Ошибка извлечения текста из PDF: {str(e)}"

    @classmethod
    def _extract_docx(cls, file_path: str) -> Tuple[bool, str]:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            
            full_text = "\n".join(paragraphs).strip()
            if not full_text:
                return False, "Документ Word пуст или не содержит текста."
            return True, full_text
        except Exception as e:
            logger.warning(f"python-docx error: {e}")
            return False, f"Ошибка извлечения текста из DOCX: {str(e)}"

    @classmethod
    def _extract_pptx(cls, file_path: str) -> Tuple[bool, str]:
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            slides_text = []
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_parts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            t = paragraph.text.strip()
                            if t:
                                slide_parts.append(t)
                if slide_parts:
                    slides_text.append(f"--- Слайд {slide_num} ---\n" + "\n".join(slide_parts))
            
            full_text = "\n\n".join(slides_text).strip()
            if not full_text:
                return False, "Презентация не содержит текста."
            return True, full_text
        except Exception as e:
            logger.warning(f"python-pptx error: {e}")
            return False, f"Ошибка извлечения текста из PPTX: {str(e)}"
